"""Headless smoke test for the Spikeless UI. QT_QPA_PLATFORM=offscreen.

Drives a synthetic GINA export through the four steps the way the mouse would — spikes
excluded on load and one put back, the background three ways, regions found, drawn, moved
and removed, a shared edge, decay, method names — then the results table, the report in every
format, the figure, a session saved and opened again, the Options window, undo / redo; the
plot's gestures (zoom to a span, an elastic pan and wheel, the reading line, Draw region by two
clicks, a stretch's edge), run names, selecting none, finding again, % of total, a legend
outside, a .json session, the method's conditions, a new run's settings, an exported trace
pasted back; and the reference values on the example runs in C:\\Code\\Spikeless\\data when
that folder is there.
Run it after touching anything:  python misc/smoke_test.py
"""
import json
import os
import tempfile
import zipfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import numpy as np  # noqa: E402
from matplotlib.backend_bases import MouseEvent  # noqa: E402
from PySide6.QtCore import QEvent, QLocale, QPoint, Qt  # noqa: E402
from PySide6.QtTest import QTest  # noqa: E402
from PySide6.QtWidgets import (QApplication, QCheckBox, QMessageBox,  # noqa: E402
                               QTableWidget)

from spikeless import figure, io_gina, options, report, session  # noqa: E402
from spikeless.editors import GREY, WHITE  # noqa: E402
from spikeless.adjust import decay_on, lookup_half_life_s  # noqa: E402

TMP = Path(tempfile.gettempdir())
options.OPTIONS_FILE = TMP / "_spikeless_smoke_options.json"   # the shipped defaults,
options.OPTIONS_FILE.unlink(missing_ok=True)                    # not the user's file
QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.Yes)
QMessageBox.warning = staticmethod(lambda *a, **k: None)

import spikeless.app as appmod  # noqa: E402
from spikeless.app import MainWindow, _rows_of  # noqa: E402

# ---- a GINA export: 7 cps background, two peaks, three single-point spikes (one on a peak) ----
n = 900
t = np.arange(n, dtype=float)
lam = 7 + 1500 * np.exp(-0.5 * ((t - 480) / 7) ** 2) + 120 * np.exp(-0.5 * ((t - 200) / 8) ** 2)
y = np.random.default_rng(7).poisson(lam).astype(float)   # seeded: the assertions are fixed
spike_at = [100, 476, 700]                                # 476: on the big peak's rising edge
y[spike_at] += 4000
lines = ["Interval Values Smoke Run, 07/07/2026 14:17", '"Time, s"\t"Taux comptage du canal A, cps"']
lines += ['"%02d\\"00"\t"%s"' % (i, ("%.3f" % v).replace(".", ",")) if i < 60 else
          '"%02d\'%02d"\t"%s"' % (i // 60, i % 60, ("%.3f" % v).replace(".", ","))
          for i, v in enumerate(y)]
src = TMP / "spikeless_smoke.txt"
src.write_bytes(("\n".join(lines) + "\n").encode("cp1252"))

app = QApplication([])
w = MainWindow()
assert not w.runs, "starts empty"


def last_log():
    return w.log_view.toPlainText().rstrip().splitlines()[-1]


# ---- ① load: spikes excluded at once, the "Interval Values" prefix dropped ----
w._load(str(src))
run = w.sel
assert len(w.runs) == 1 and run.name == "Smoke Run", run.name
assert run.spike_list() == spike_at, run.spike_list()
assert "3 spike(s) excluded" in w.log_view.toPlainText()
assert len(run.sha256) == 64
assert w.spike_list.count() == 3 and w.step["spikes"]._btn.text() == "Spikes · 3 excluded"

# ---- ④ Find peaks: both peaks, % of total with the outside row adding to 100 ----
w._find()
res = w.results[run.id]
assert len(res.rows) == 2, [(r.start, r.end) for r in res.rows]
small, big = res.rows
assert abs(small.rt - 200 / 60) < 0.05 and abs(big.rt - 8.0) < 0.05, (small.rt, big.rt)
assert abs(sum(r.pct for r in res.rows) + 100 * res.outside / res.total - 100) < 1e-9
true_big = 1500 * 7 * np.sqrt(2 * np.pi)
assert abs(big.net - true_big) < 3 * big.unc + 0.02 * true_big, (big.net, true_big)
assert big.spikes == 1 and big.net_with_spikes > big.net + 2000, big
assert abs(res.background - 7) < 0.5, res.background
assert w.table.rowCount() == 4, w.table.rowCount()     # 2 regions + outside + total
heads = [w.table.horizontalHeaderItem(c).text() for c in range(w.table.columnCount())]
assert heads[:1] == ["Region"] and "% of total" in heads and "Net counts" in heads, heads

# ---- ② a spike put back is back; the list's tick is the switch; undo ----
it = w.spike_list.item(1)                                  # the one on the big peak
it.setCheckState(Qt.Unchecked)
assert run.kept == [476] and run.trace()[476] == y[476], "unticking did not put it back"
assert w.results[run.id].rows[1].net > big.net + 2000, "the area did not take the spike back"
w._undo_step()
run = w.sel
assert run.kept == [] and last_log() .endswith("Undone: keep the spike at 7.93 min.")
w._redo_step()
assert w.sel.kept == [476]
w._undo_step()
run = w.sel

# ---- ② the threshold: higher finds fewer, undo restores ----
w.sens.setValue(30)                                         # 15 σ
assert run.spike_params.n_sigma == 15 and len(run.spike_list()) == 3   # 4000 cps: still spikes
w._undo_step()
run = w.sel
assert run.spike_params.n_sigma == 5

# ---- ③ background: a stretch, a drift, back to auto ----
w._stretch(0.5, 1.5)
run = w.sel
sel = (run.x >= 0.5) & (run.x <= 1.5)
assert run.bg_mode == "stretch" and abs(w.results[run.id].background
                                        - run.trace()[sel].mean()) < 1e-9
w.bg_radio["drift"].setChecked(True)
assert run.bg_mode == "drift" and "SNIP" in w.results[run.id].background_how
w.bg_radio["auto"].setChecked(True)
assert run.bg_mode == "auto"

# ---- ④ drawing: snapped, clipped against neighbours, n.d. on noise; edges; Del ----
r0, r1 = run.regions
w._drawn(11.0, 11.6, False)                                 # noise only
k = w.sel_region
row = w.results[run.id].rows[k]
assert run.regions[k].how == "drawn" and "n.d." in row.flags and row.pct is None, row
w._delete_regions([w.sel_region])
assert len(run.regions) == 2
w._drawn(r1.start - 0.3, r1.start + 0.2, True)              # overlaps the big one: cut at its start
k = w.sel_region
assert abs(run.regions[k].end - r1.start) < 1e-9, (run.regions[k], r1)
# that edge is now shared: moving it moves both
shared = run.regions[k].end
w._edge_moved(k, "end", shared - 0.05, True)
assert abs(run.regions[k].end - (shared - 0.05)) < 1e-9
assert abs(run.regions[k + 1].start - (shared - 0.05)) < 1e-9
assert run.regions[k + 1].how == "moved", run.regions[k + 1].how
w._undo_step()
w._undo_step()                                              # the edge, then the drawn region
run = w.sel
assert len(run.regions) == 2 and run.regions[1].how == "auto"
# ← → nudge the last-touched edge one point
w.canvas._edge = (1, "end")
end0 = run.regions[1].end
w.canvas.keyPressEvent(type("K", (), {"key": lambda self: Qt.Key_Right})())
assert abs(w.sel.regions[1].end - end0 - 1 / 60) < 1e-9
w._undo_step()
run = w.sel
# a region's line under it, its note, its name typed in the table
w.sel_region = 1
w.refresh()
w.region_panel()                                            # Edit…: the region on its own
rp = w.region_win
rp.refresh()
assert rp.form.isEnabled() and "peak 2" in rp.head.text(), rp.head.text()
rp.under.setCurrentIndex(1)
rp.under.activated.emit(1)
assert run.regions[1].under == "line"
w._undo_step()
run = w.sel
w.table.item(1, 0).setText("product")
assert run.regions[1].name == "product" and w.results[run.id].rows[1].name == "product"

# ---- methods name regions; decay per isotope ----
options.PREFS["methods"] = [["M", "free", 3.0, 3.6], ["M", "product", 7.5, 8.5]]
run.method = "M"
w.refresh()
assert w.results[run.id].rows[0].name == "free" and w.results[run.id].rows[0].named
base = run.trace()
options.PREFS["isotopes"][6][2] = True                      # ⁶⁸Ga: decay-corrected
options.set_half_lives(options.PREFS["isotopes"])
run.isotope = "⁶⁸Ga"
w.refresh()
k = run.trace()[600] / base[600]                            # 10 min after the run start
assert run.decayed() and abs(k - 2 ** (10 / (1.12782 * 60))) < 1e-9, k
assert "decay-corrected" in w.summary.text()
options.PREFS["isotopes"][6][2] = False
options.set_half_lives(options.PREFS["isotopes"])
w.refresh()
assert not run.decayed()

# ---- the report: preview and every format carry the same numbers ----
rep = report.html(w)
assert "Smoke Run" in rep and "Spikes: 3 excluded" in rep and "Hampel" in rep, rep[:800]
assert "with its excluded spike(s) kept" in rep, "the spike-kept area is missing"
w._report()
rw = w.report_win
assert "Smoke Run" in rw.view.toPlainText()
for ext in (".pdf", ".odt", ".xlsx", ".md", ".html"):
    out = TMP / f"spikeless_smoke_report{ext}"
    rw.write(out)
    assert out.stat().st_size > 1000, (ext, out.stat().st_size)
rt = f"{w.results[run.id].rows[1].rt:.2f}"
assert rt in (TMP / "spikeless_smoke_report.md").read_text(encoding="utf-8")
assert (TMP / "spikeless_smoke_report_plot.png").exists(), "the markdown's figure is not beside it"
assert any(nm.startswith("Pictures/") for nm in
           zipfile.ZipFile(TMP / "spikeless_smoke_report.odt").namelist()), "no figure in the .odt"
assert "data:image/png;base64" in (TMP / "spikeless_smoke_report.html").read_text("utf-8")
from openpyxl import load_workbook  # noqa: E402
wb = load_workbook(TMP / "spikeless_smoke_report.xlsx")
assert {"Report", "Results", "Spikes", "Figure"} <= set(wb.sheetnames), wb.sheetnames
rs = list(wb["Results"].values)
assert rs[0][:2] == ("Run", "Region") and rs[2][1] == "product", rs[:3]
assert rs[2][rs[0].index("Rt (min)")] == float(rt), rs[2]
assert len(list(wb["Spikes"].values)) == 4                  # header + 3 spikes
# a section unticked leaves the report and is saved; a column unticked leaves the table too
tree = rw.tree
results_item = next(tree.topLevelItem(i) for i in range(tree.topLevelItemCount())
                    if tree.topLevelItem(i).text(0) == "Results")
unc = next(results_item.child(i) for i in range(results_item.childCount())
           if results_item.child(i).text(0) == "± counts")
unc.setCheckState(0, Qt.Unchecked)
assert "± counts" not in report.html(w)
assert "± counts" in [w.table.horizontalHeaderItem(c).text()        # the table: its own
                      for c in range(w.table.columnCount())]
assert json.loads(options.OPTIONS_FILE.read_text("utf-8"))["report_cols"]["unc"] is False
unc.setCheckState(0, Qt.Checked)

# ---- the figure: its size in mm, presets, copy ----
fig = figure.figure(w)
assert all(abs(a - b) < 0.5 for a, b in zip(fig.plot_area_mm, (80, 40))), fig.plot_area_mm
w._figure()
fw = w.figure_win
fw._preset(4)                                               # Slide
assert w.plot_opts.plot_w_mm == 160 and w.plot_opts.x.tick_font_size == 14
assert abs(figure.figure(w).plot_area_mm[0] - 160) < 0.5
figure.copy_figure(w)
assert not QApplication.clipboard().image().isNull(), "nothing on the clipboard"
texts = [t.get_text() for t in figure.figure(w).axes[0].texts]
assert any("product" in s for s in texts), texts            # named regions labelled
next(c for c in fw.side.widget().findChildren(QCheckBox)
     if c.text() == "Shaded").setChecked(False)
assert not figure.figure(w).axes[0].collections, "regions still shaded"
w.plot_opts.show_regions = True

# ---- a session comes back as it was saved ----
ses = TMP / "spikeless_smoke.spikeless"
ses.write_text(session.dumps(w.runs, w.plot_opts), encoding="utf-8")
before = [(r.start, r.end, r.name, r.how) for r in w.sel.regions]
w._remove_run()
assert not w.runs
w._load(str(src))                                           # something for it to replace
w._open_session(str(ses))
assert [(r.start, r.end, r.name, r.how) for r in w.sel.regions] == before
assert w.sel.method == "M" and w.plot_opts.plot_w_mm == 160 and ses.name in w.windowTitle()
w._undo_step()
assert w.sel.regions == [] and w.sel.name == "Smoke Run", "undo did not give back what was replaced"
w._redo_step()

# ---- the Options window: saved as changed, followed at once, reset ----
w._options("Isotopes")
ow = w.options_win
assert ow.topics.currentItem().text() == "Isotopes"
tab = ow.pages.currentWidget().findChildren(QTableWidget)[0]
tab.item(0, 1).setText("7")
assert lookup_half_life_s(tab.item(0, 0).text()) == 7 * 3600
tab.item(9, 2).setCheckState(Qt.Checked)                    # ¹⁷⁷Lu decay-corrected
assert decay_on("¹⁷⁷Lu")
w._options("Methods")
mt = ow.pages.currentWidget().findChildren(QTableWidget)[0]
assert mt.rowCount() == 2 and mt.item(1, 1).text() == "product"
ow._apply(None)                                             # Reset to defaults
QApplication.sendPostedEvents(None, QEvent.DeferredDelete)
assert lookup_half_life_s("⁹⁹ᵐTc") == 6.00718 * 3600 and not decay_on("¹⁷⁷Lu")
assert options.PREFS["methods"] == [] and w.results[w.sel.id].rows[0].name.startswith("peak 1")

assert QLocale().decimalPoint() == ".", "a decimal comma in the fields"

# ---- the plot: zoom to a span, an elastic pan and wheel, the reading line ----
w.show()
app.processEvents()
cv = w.canvas
cv.anim_ms = 0                                              # no sliding: at once
cv.draw()


def at(t, frac=0.5):
    y0, y1 = cv.ax.get_ylim()
    return cv.ax.transData.transform((t, y0 + frac * (y1 - y0)))


def ev(name, p, button=None, step=0):
    return MouseEvent(name, cv, p[0], p[1], button=button, step=step)


def drag(t0, t1, button=1):
    p0, p1 = at(t0), at(t1)
    cv._on_press(ev("button_press_event", p0, button))
    cv._on_motion(ev("motion_notify_event", p1))
    cv._on_release(ev("button_release_event", p1, button))


def click(t, button=1):
    p = at(t)
    cv._on_press(ev("button_press_event", p, button))
    cv._on_release(ev("button_release_event", p, button))


drag(2.0, 4.0)                                              # left drag: zoom to the span
assert cv.xlim and abs(cv.xlim[0] - 2) < 0.05 and abs(cv.xlim[1] - 4) < 0.05, cv.xlim
assert cv.ylim is None                                      # the height fitted to it
drag(2.2, 5.5, button=3)                                    # right drag past the start…
assert abs(cv.xlim[0]) < 1e-9 and abs(cv.xlim[1] - cv.xlim[0] - 2) < 0.05, cv.xlim  # …back
click(1.2)                                                  # a reading line, centred on it
assert abs(cv.cursor_t - 1.2) < 0.05 and abs(sum(cv.xlim) / 2 - cv.cursor_t) < 0.05, cv.xlim
reading = next(t.get_text() for t in cv.ax.texts if "point" in t.get_text())
assert "min · point" in reading and "/ 900" in reading and "background" in reading, reading
for _ in range(12):                                         # the wheel out past the run…
    cv._on_scroll(ev("scroll_event", at(1.0), button="down", step=-1))
cv._settle_later.stop()
cv._settle()
assert cv.xlim is None and cv.ylim is None                  # …springs back to all of it
it = w.spike_list.item(0)                                   # a click on a spike: zoomed on it
w._spike_focus(it)
t0 = w.sel.x[it.data(Qt.UserRole)]
assert cv.xlim[0] < t0 < cv.xlim[1] and cv.cursor_t == t0, (cv.xlim, t0)
cv.reset_view()

# ---- Draw region: pushed in, two clicks; a right click cancels ----
n = len(w.sel.regions)
w.reg_draw.setChecked(True)
assert cv.mode == "region"
click(11.0)
assert cv.anchor is not None and w.reg_draw.isChecked()
click(11.6)
assert cv.mode is None and not w.reg_draw.isChecked() and len(w.sel.regions) == n + 1
w._undo_step()
w.reg_draw.setChecked(True)
click(5.0, button=3)
assert cv.mode is None and not w.reg_draw.isChecked() and len(w.sel.regions) == n

# ---- the background stretch: shaded, its edge dragged, its ends typed ----
w.bg_radio["stretch"].setChecked(True)
run = w.sel
assert run.bg_mode == "stretch" and sorted(run.bg_stretch) == [0.5, 1.5], run.bg_stretch
drag(1.5, 2.0)                                              # its end's edge
assert abs(w.sel.bg_stretch[1] - 2.0) < 0.02 and abs(w.bg_to.value() - 2.0) < 0.02
w.bg_from.setValue(0.2)
assert abs(w.sel.bg_stretch[0] - 0.2) < 1e-9
w.bg_radio["auto"].setChecked(True)

# ---- run names: an alias typed in the list, a legend of its own, Run names ▤ ----
run = w.sel
w.run_list.item(0).setText("Alias")
assert (run.name, run.source, run.legend()) == ("Alias", "Smoke Run", "Alias")
w.relabel(run, "On the figure")
assert run.legend() == "On the figure" and run.name == "Alias"
w.names_window()
nw = w.names_win
nw.refresh()
assert nw.t.isColumnHidden(0) and nw.t.item(0, 0).text() == "Smoke Run"
assert nw.t.item(0, 1).foreground().color() == WHITE and nw.t.item(0, 0).foreground().color() == GREY
nw.t.item(0, 2).setText("")                                 # cleared: follows the name again
app.processEvents()
assert run.legend() == "Alias" and nw.t.item(0, 2).foreground().color() == GREY
nw.t.item(0, 1).setText("")                                 # cleared: the original again
app.processEvents()
assert run.name == "Smoke Run"

# ---- two runs: a legend; outside, the figure grows for it ----
w._load(str(src))
o = w.plot_opts
o.legend_loc = "in upper right"
w_in = figure.figure(w).get_figwidth()
o.legend_loc = "out right top"
f_out = figure.figure(w)
assert f_out.get_figwidth() > w_in and abs(f_out.plot_area_mm[0] - o.plot_w_mm) < 0.5
assert [t.get_text() for t in f_out.axes[0].get_legend().get_texts()] == ["Smoke Run"] * 2
o.legend_loc = "auto"

# ---- selecting none: the steps wait; a click below the runs does it ----
w._select(None)
assert w.sel is None and w.step["spikes"]._btn.text().endswith("select a run")
assert "Select a run" in w.summary.text() and not w.body["peaks"].isEnabled()
assert w.ticks["spikes"].isEnabled()                        # its Show ticks stay at hand
w._run_picked(0)
assert w.sel is w.runs[0]
vp = w.run_list.viewport()
QTest.mouseClick(vp, Qt.LeftButton, pos=QPoint(10, vp.height() - 3))
assert w.sel is None, "a click below the runs did not select none"
w._run_picked(0)

# ---- % of total: each run's net counts, per min ----
w.y_mode.setCurrentIndex(2)
w.y_mode.activated.emit(2)
assert w.plot_opts.y_mode == "pct_total"
assert figure.figure(w).axes[0].get_ylabel() == "% of total"
x_, v_ = w.canvas._curves[0]
assert abs(v_.sum() - 100) < 1, v_.sum()                   # each point's share: 100 % in all
w.y_mode.setCurrentIndex(0)
w.y_mode.activated.emit(0)

# ---- finding again: the threshold moves, a region with a name of its own stays ----
w._find()
run = w.sel
n = len(run.regions)
w.set_region(0, "name", "kept")
w.find_sens.setValue(12)                                    # 6 σ
assert run.find_params.n_sigma == 6 and run.regions[0].name == "kept" and len(run.regions) == n
assert "6 σ above" in " ".join(report.processing(run, w.results[run.id]))

# ---- the regions in full (⤢): a note, a label off the figure; ⧉; Del on several ----
w.regions_window()
rg = w.regions_win
rg.refresh()
assert rg.t.rowCount() == len(run.regions)
last = rg.t.columnCount() - 1
rg.t.item(0, last).setText("checked by hand")
assert w.sel.regions[0].note == "checked by hand"
app.processEvents()
rg.t.item(0, last - 1).setCheckState(Qt.Unchecked)
assert w.sel.regions[0].label is False
w._copy_results()
assert QApplication.clipboard().text().startswith("Region")
w.table.selectAll()
w._delete_regions(_rows_of(w.table))
assert w.sel.regions == []
w._undo_step()
assert len(w.sel.regions) == n

# ---- the method's conditions in the report ----
options.PREFS["method_info"] = [["M", "Column", "C18, 150 × 4.6 mm"], ["M", "Gradient", ""]]
w.sel.method = "M"
rep = report.html(w)
assert "C18, 150 × 4.6 mm" in rep and "Gradient" not in rep

# ---- a new run starts with the options' settings ----
options.PREFS["spikes"]["n_sigma"] = 15.0
options.PREFS["find"]["n_sigma"] = 8.0
w._load(str(src))
assert w.sel.spike_params.n_sigma == 15 and w.sel.find_params.n_sigma == 8
options.load_prefs(got={})
w._options("Chromatogram")
assert w.options_win.topics.currentItem().text() == "Chromatogram"

# ---- an exported trace, pasted back: the same numbers; as recorded too ----
src_run = w.runs[0]
out = TMP / "spikeless_smoke_export.txt"
appmod.QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (str(out), ""))
w._export_trace(src_run)
text = out.read_text("cp1252")
assert "," not in text.split("\n", 2)[2], "a decimal comma in the export"
QApplication.clipboard().setText(text)
n = len(w.runs)
w._paste()
assert len(w.runs) == n + 1 and np.allclose(w.runs[-1].y, src_run.trace(), atol=1e-3)
w._export_trace(src_run, recorded=True)
assert np.allclose(io_gina.load(out)[0].y, src_run.y, atol=1e-3)

# ---- a session as .json: saved, opened ----
js = TMP / "spikeless_smoke.json"
appmod.QFileDialog.getSaveFileName = staticmethod(
    lambda *a, **k: (str(TMP / "spikeless_smoke"), "JSON — the same session (*.json)"))
w._save_session()
assert json.loads(js.read_text("utf-8"))["app"] == "Spikeless"
n = len(w.runs)
w._load_paths([str(js)])
assert len(w.runs) == n and js.name in w.windowTitle()

# ---- what the panels show; the regions list; the results' own columns; Copy the screen ----
assert not w.run_rows.isRowVisible(w.f_iso) and not w.run_rows.isRowVisible(w.f_method)
assert w.run_rows.isRowVisible(w.f_mol)                     # no isotope, no method at first
options.PREFS["run_fields"]["isotope"] = True
w.refresh()
assert w.run_rows.isRowVisible(w.f_iso)
options.PREFS["run_fields"]["isotope"] = False
if not w.sel.regions:
    w._find()
w.reg_list.item(0, 0).setText("renamed in the list")
assert w.sel.regions[0].name == "renamed in the list"
assert w.results[w.sel.id].rows[0].name == "renamed in the list"
options.PREFS["table_cols"]["height"] = True
w.refresh()
assert "Height (cps)" in [w.table.horizontalHeaderItem(c).text()
                          for c in range(w.table.columnCount())]
assert "Height (cps)" not in report.html(w)                 # the report keeps its own
options.PREFS["table_cols"]["height"] = False
w.view["stretch"] = False                                   # Show › Stretch off: not shaded
w.sel.bg_mode, w.sel.bg_stretch = "stretch", [0.5, 1.5]
w.refresh()
assert not any("stretch" in t.get_text() for t in w.canvas.ax.texts)
w.view["stretch"] = True
w.refresh()
assert any("stretch" in t.get_text() for t in w.canvas.ax.texts)
QApplication.clipboard().clear()
w._copy_screen()
img = QApplication.clipboard().image()
assert not img.isNull() and img.width() == w.canvas.width(), img.size()
x0 = w.canvas.ax.get_position().x0                          # the plot's left edge stays put
w.canvas.focus(7.5, 8.5)                                    # as the height (its labels) changes
w.canvas.redraw()
assert w.canvas.ax.get_position().x0 == x0
w.canvas.reset_view()

# ---- the example runs: reference values (devlog) ----
data = Path(__file__).resolve().parents[2] / "data"
if data.is_dir():
    def ref(stem):
        w._load(str(next(data.glob(f"Interval_Values_{stem}*.txt"))))
        w._find()
        return w.sel, w.results[w.sel.id]
    r, res = ref("HOPO-1")
    assert r.spike_list() == [587, 700] and len(res.rows) == 2, (r.spike_list(), res.rows)
    assert [round(x.pct, 1) for x in res.rows] == [14.3, 85.7], [x.pct for x in res.rows]
    r, res = ref("Imathera2")
    assert r.spike_list() == [778, 786], r.spike_list()
    imp = next(x for x in res.rows if abs(x.rt - 2.55) < 0.05)
    prod = next(x for x in res.rows if abs(x.rt - 7.90) < 0.05)
    assert abs(imp.pct - 2.2) < 0.15 and abs(prod.pct - 97.7) < 0.15, (imp.pct, prod.pct)
    r, res = ref("TestOnduleur")
    assert r.spike_list() == [1314, 4696, 5660] and res.rows == [] and not res.active

w.close()
for p in [src, ses, out, js, options.OPTIONS_FILE, TMP / "spikeless_smoke_report_plot.png"] + \
        [TMP / f"spikeless_smoke_report{e}" for e in (".pdf", ".odt", ".xlsx", ".md", ".html")]:
    p.unlink(missing_ok=True)
print("spikeless smoke OK")
