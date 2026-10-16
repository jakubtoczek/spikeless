"""Spikeless - the report: what was loaded, what was done to it, the results, the figure.

On the ticked runs or the selected one (SCOPES). Several runs: a summary first — a line per
run, every run's results in one table — then what was done to each, then their details, what
they share once and what differs a column each; or each run in full, one after the other.

One HTML text, built from the session as it stands, is both the preview and every format:
PDF and .odt are Qt's print of it, .md Qt's markdown of it, .html it with the figure inlined;
the workbook takes the same tables, a sheet each. The Report window picks the format, the
sections and the results' columns (the table under the plot has its own: Options › Results),
saved with the options the moment they change.
"""

from __future__ import annotations

import base64
import datetime as _dt
from html import escape as esc
from io import BytesIO
from pathlib import Path

from PySide6.QtCore import QBuffer, QByteArray, QMimeData, Qt, QTimer, QUrl
from PySide6.QtGui import QKeySequence, QPageSize, QPdfWriter, QShortcut, QTextDocument, \
    QTextDocumentWriter
from PySide6.QtWidgets import (
    QApplication, QComboBox, QFileDialog, QHBoxLayout, QLabel, QMainWindow, QMessageBox,
    QPushButton, QTextBrowser, QTreeWidget, QTreeWidgetItem, QVBoxLayout, QWidget,
)

from . import __version__
from .options import PREFS, method_info, save_prefs


def _n(v, fmt):
    return "" if v is None else format(v, fmt)


def _pct(row, key):
    v = getattr(row, key)
    return "n.d." if "n.d." in row.flags else _n(v, ".1f")


# The results' columns (key, header, value(row)): the same spec drives the table under the
# plot, the report and the workbook, so what is shown is what is written.
COLUMNS = [
    ("name", "Region", lambda r: r.name),
    ("start", "Start (min)", lambda r: f"{r.start:.2f}"),
    ("end", "End (min)", lambda r: f"{r.end:.2f}"),
    ("rt", "Rt (min)", lambda r: f"{r.rt:.2f}"),
    ("net", "Net counts", lambda r: f"{r.net:.0f}"),
    ("unc", "± counts", lambda r: f"{r.unc:.0f}"),
    ("pct", "% of total", lambda r: _pct(r, "pct")),
    ("pct_reg", "% of regions", lambda r: _pct(r, "pct_reg")),
    ("height", "Height (cps)", lambda r: f"{r.height:.0f}"),
    ("w_half", "w½ (min)", lambda r: _n(r.w_half, ".3f")),
    ("sym", "Symmetry", lambda r: _n(r.sym, ".2f")),
    ("rs", "Resolution", lambda r: _n(r.rs, ".1f")),
    ("sn", "S/N", lambda r: _n(r.sn, ".0f")),
    ("flags", "Flags", lambda r: ", ".join(r.flags)),
]
HELP = {"start": "Where the region begins, in min",
        "end": "Where it ends, in min",
        "rt": "Retention time: where it is highest above what lies under it, in min",
        "net": "Its counts above the background (or a straight line between its ends): the "
               "count rate × time, by trapezoids",
        "unc": "The 1 σ uncertainty of the net counts from counting statistics. Net counts "
               "N = G − B: G every count recorded across the region, B the background's counts "
               "under it. Each is a Poisson count, whose variance is the count itself, so "
               "σ = √(G + B). The true value lies within ± 1 σ about 2 times in 3, within "
               "± 2 σ 19 times in 20. Divided by N, the relative uncertainty: 10 % at the "
               "quantification limit (< LOQ). B is taken as if counted over the region alone "
               "— it is the mean of far more points, so this errs on the safe side — and "
               "nothing but counting is in it (not the injection, the detector's efficiency…)",
        "pct": "% of all the net counts of the run, inside and outside the regions: the "
               "radiochemical purity as Ph. Eur. defines it (of the total radioactivity). The "
               "regions and Outside the regions add up to 100 %",
        "pct_reg": "% of the regions' net counts only (those below the detection limit left "
                   "out): what a purity over the integrated peaks alone would say",
        "height": "Its highest point above what lies under it, in cps",
        "w_half": "Width at half height, in min",
        "sym": "Symmetry factor w₀.₀₅ / 2d (Ph. Eur. 2.2.46): 1 symmetric, above 1 tailing",
        "rs": "Resolution from the region before, 1.18 (t₂ − t₁) / (w½₁ + w½₂) (Ph. Eur. "
              "2.2.46): 1.5 and above, separated to the baseline",
        "sn": "Signal-to-noise 2H / h, h the range of the background around the region "
              "(Ph. Eur. 2.2.46)",
        "flags": "drawn / moved / copied: placed by hand · < LOQ: below the quantification "
                 "limit · n.d.: below the detection limit (Currie) · spike excluded: a spike "
                 "inside the region was excluded"}
SECTIONS = [("plot", "Figure", "As Copy / Export make it (Figure window)"),
            ("runs", "Runs", "File and its SHA-256, run date, isotope, molecule, method and "
                             "its conditions (Options › Methods), the run's conditions"),
            ("processing", "Processing", "Spikes excluded, background, how the regions were "
                                         "found, decay correction"),
            ("spikes", "Spikes", "Each spike: time, as recorded, as bridged, in which region"),
            ("results", "Results", "A row per region, outside the regions, total — in the "
                                   "columns ticked")]
SCOPES = [("summary", "Ticked runs: a summary, each run, their details"),
          ("each", "Ticked runs, each in full, one after the other"),
          ("selected", "The selected run only")]
SUMMARY = ("Run", "Spikes excluded", "Background (cps)", "Net counts", "Regions", "Decay")
FORMATS = [(".pdf", "PDF"), (".odt", "Document (.odt) — opens in Word and LibreOffice"),
           (".xlsx", "Excel workbook — a sheet per table"), (".md", "Markdown text"),
           (".html", "Web page")]
PLOT = "plot.png"                      # the figure's name inside the report
LEGEND = ("Net counts above the background (or a straight line between a region's ends), "
          "± 1 σ counting statistics; % of total = of all the net counts of the run. "
          "Flags: drawn / moved / copied = placed by hand; < LOQ = below the quantification "
          "limit, n.d. = below the detection limit (Currie 1968); spike excluded = a spike "
          "inside the region was excluded.")


def columns(which="report_cols"):
    """The Region column, always (it names the row), then the ones ticked: for the report
    ("report_cols", the Report window) or the table under the plot ("table_cols", Options)."""
    return [(k, h, f) for k, h, f in COLUMNS if k == "name" or PREFS[which].get(k)]


def rows(res, which="report_cols"):
    """The ticked columns of each region, then outside the regions and total."""
    cols = columns(which)
    body = [[f(r) for _, _, f in cols] for r in res.rows]
    extra = []
    for name, net, pct in (("Outside the regions", res.outside,
                            100 * res.outside / res.total if res.active else None),
                           ("Total", res.total, 100.0 if res.active else None)):
        extra.append([name if k == "name" else f"{net:.0f}" if k == "net" else
                      _n(pct, ".1f") if k == "pct" else "" for k, _, _ in cols])
    return [h for _, h, _ in cols], body, extra


def spike_rows(run, res):
    """Each flagged point: time, as recorded, as bridged, region, excluded or kept."""
    from .spikes import remove
    clean = remove(run.y, run.excluded())          # before any decay correction, like y
    names = [(r.start, r.end, row.name) for r, row in zip(run.regions, res.rows)]
    out = []
    for i in run.spike_list():
        t = float(run.x[i])
        where = next((n for a, b, n in names if a <= t <= b), "")
        out.append([f"{t:.2f}", f"{run.y[i]:.0f}",
                    f"{clean[i]:.0f}" if i not in run.kept else "—", where,
                    "kept (by hand)" if i in run.kept else "excluded"])
    return out


def processing(run, res) -> list[str]:
    """What was done to the run, a sentence each."""
    p = run.spike_params
    n = int(run.excluded().sum())
    out = [f"Spikes: {n} excluded"
           + (f", {len(run.kept)} flagged and kept by hand" if run.kept else "")
           + f" — Hampel filter, {p.window}-point window, {p.n_sigma:g} σ, at most "
           f"{p.max_width} point(s) wide, clear of both neighbours; each excluded point "
           "bridged linearly between its neighbours."]
    for row in res.rows:
        if row.spikes:
            d = row.net_with_spikes - row.net
            out.append(f"{row.name}: {row.net:.0f} counts; with its excluded spike(s) kept, "
                       f"{row.net_with_spikes:.0f} ({100 * d / row.net:+.1f} %)."
                       if row.net else f"{row.name}: spike(s) excluded.")
    out.append(f"Background: {res.background:.2f} cps — {res.background_how}.")
    if run.regions:
        hows = [r.how for r in run.regions]
        found, moved = hows.count("auto") + hows.count("moved"), hows.count("moved")
        f, parts = run.find_params, []
        if found:
            parts.append(f"{found} found (apexes of the {f.smooth}-point moving mean "
                         f"{f.n_sigma:g} σ above their surroundings, at least {f.min_gap_s:g} s "
                         f"apart, out to their feet at {f.foot_sigma:g} σ; kept above the "
                         "detection limit)" + (f", {moved} of them moved by hand" if moved
                                               else ""))
        if len(hows) > found:
            parts.append(f"{len(hows) - found} placed by hand")
        out.append("Regions: " + "; ".join(parts) + ".")
    out.append(f"Decay: corrected to the run start ({run.isotope}, half-life "
               f"{run.half_life_s() / 3600:g} h)." if run.decayed() else
               "Decay: not corrected.")
    if not res.active:
        out.append("No activity above background: the run's net counts are below its "
                   "detection limit, so no % is given.")
    return out


def run_info(run) -> list[tuple[str, str]]:
    """The run's metadata: its original name if renamed, the file, the method and the
    method's conditions (Options › Methods), the run's own conditions."""
    return [(k, v) for k, v in (
        ("Original name", run.source if run.source != run.name else ""),
        ("File", run.file), ("SHA-256", run.sha256), ("Run", run.run_datetime),
        ("Points", f"{len(run.y)}, one every {run.dt_s:g} s, {run.x[-1]:.2f} min"),
        ("Isotope", run.isotope), ("Molecule", run.molecule), ("Method", run.method),
        *method_info(run.method),
        ("Conditions", "; ".join(c for c in run.conditions.splitlines() if c.strip())))
        if v]


def scope() -> str:
    return PREFS["report_scope"] if PREFS["report_scope"] in dict(SCOPES) else SCOPES[0][0]


def report_runs(main) -> list:
    """The runs the report is on: the selected one (else the first ticked), or every ticked
    one, in the list's order."""
    shown = [r for r in main.runs if r.shown]
    if scope() == "selected":
        run = main.sel if main.sel is not None else next(iter(shown), None)
        return [run] if run is not None else []
    return shown


def summary_rows(main, runs) -> list[list]:
    """A line per run (SUMMARY's columns)."""
    out = []
    for r in runs:
        res = main.results[r.id]
        out.append([r.name, int(r.excluded().sum()), f"{res.background:.1f}",
                    f"{res.total:.0f}", len(r.regions),
                    f"corrected ({r.isotope})" if r.decayed() else "not corrected"])
    return out


def all_results(main, runs):
    """Every run's results in one table, a Run column first: (head, body)."""
    body = []
    for r in runs:
        head, rows_, extra = rows(main.results[r.id])
        body += [[r.name] + row for row in rows_ + extra]
    return ["Run"] + [h for _, h, _ in columns()], body


def info_split(runs):
    """The runs' details (run_info): (what all share, as rows; the fields that differ; a row
    per run of those)."""
    infos = [dict(run_info(r)) for r in runs]
    keys = list(dict.fromkeys(k for i in infos for k in i))
    same = [(k, infos[0][k]) for k in keys
            if all(k in i and i[k] == infos[0][k] for i in infos)]
    diff = [k for k in keys if k not in dict(same)]
    return same, diff, [[r.name] + [i.get(k, "") for k in diff] for r, i in zip(runs, infos)]


def _table(head, body, border=1) -> str:
    h = "<tr>" + "".join(f"<th>{esc(c)}</th>" for c in head) + "</tr>" if head else ""
    return (f"<table border='{border}' cellspacing='0' cellpadding='3'>{h}"
            + "".join("<tr>" + "".join(f"<td>{esc(str(c))}</td>" for c in r) + "</tr>"
                      for r in body) + "</table>")


def html(main, plot_width_px: int | None = None) -> str:
    """The report as HTML. With plot_width_px, the figure is the image named PLOT."""
    on = PREFS["report_sections"]
    runs = report_runs(main)
    summary = scope() == "summary" and len(runs) > 1
    out = ["<h1>Radio-HPLC report</h1>",
           f"<p>{_dt.datetime.now():%Y-%m-%d %H:%M} · Spikeless {__version__}</p>"]
    if not runs:
        out.append("<p>Nothing to report: tick a run.</p>" if main.runs else
                   "<p>Nothing loaded.</p>")
    if on["plot"] and plot_width_px and runs:
        out.append(f"<p><img src='{PLOT}' width='{plot_width_px}'></p>")
    if summary:
        out.append(f"<h2>Summary — {len(runs)} runs</h2>" + _table(SUMMARY,
                                                                  summary_rows(main, runs)))
        if on["results"]:
            out.append("<h3>Results</h3>" + _table(*all_results(main, runs)))
    for run in runs:
        res = main.results[run.id]
        out.append(f"<h2>{esc(run.name)}</h2>")
        if on["runs"] and not summary:
            out.append(_table(None, run_info(run), border=0))
        if on["processing"]:
            out.append("<p>" + "<br>".join(esc(s) for s in processing(run, res)) + "</p>")
        if on["spikes"] and run.spike_list():
            out.append(_table(("Time (min)", "Recorded (cps)", "Bridged (cps)", "Region",
                               "Status"), spike_rows(run, res)))
        if on["results"]:
            head, body, extra = rows(res)
            if head and not summary:
                out.append("<h3>Results</h3>" + _table(head, body + extra))
            notes = [f"{row.name}: {r.note}" for r, row in zip(run.regions, res.rows) if r.note]
            if notes:
                out.append("<p>" + "<br>".join(esc(s) for s in notes) + "</p>")
    if summary and on["runs"]:
        same, diff, per = info_split(runs)
        out.append("<h2>The runs' details</h2>")
        if same:
            out.append("<p>The same for every run:</p>" + _table(None, same, border=0))
        if diff:
            out.append("<p>Run by run:</p>" + _table(["Run"] + diff, per))
    if on["results"] and runs:
        out.append(f"<p><small>{esc(LEGEND)}</small></p>")
    return "".join(out)


def num(v):
    """A formatted cell back to a number for Excel (so columns stay summable); else text."""
    try:
        return float(v)
    except (ValueError, TypeError):
        return v


def to_xlsx(main, path, png: bytes | None = None, width_px: int | None = None) -> None:
    """The report as a workbook: runs and processing on one sheet, the results of every run on
    one, the spikes on one, the figure on its own."""
    from openpyxl import Workbook

    on = PREFS["report_sections"]
    runs = report_runs(main)
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(["Radio-HPLC report", f"{_dt.datetime.now():%Y-%m-%d %H:%M}",
               f"Spikeless {__version__}"])
    for run in runs:
        res = main.results[run.id]
        ws.append([])
        ws.append([run.name])
        if on["runs"]:
            for k, v in run_info(run):
                ws.append([k, v])
        if on["processing"]:
            for s in processing(run, res):
                ws.append([s])
    if on["results"] and columns():
        rs = wb.create_sheet("Results")
        head = [h for _, h, _ in columns()]
        rs.append(["Run"] + head)
        for run in runs:
            _, body, extra = rows(main.results[run.id])
            for r in body + extra:
                rs.append([run.name] + [num(c) for c in r])
        rs.append([])
        rs.append([LEGEND])
    if on["spikes"]:
        sp = wb.create_sheet("Spikes")
        sp.append(["Run", "Time (min)", "Recorded (cps)", "Bridged (cps)", "Region",
                   "Status"])
        for run in runs:
            for r in spike_rows(run, main.results[run.id]):
                sp.append([run.name] + [num(c) for c in r])
    if on["plot"] and png:
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(BytesIO(png))
        img.height, img.width = round(img.height * width_px / img.width), width_px
        wb.create_sheet("Figure").add_image(img, "A1")
    wb.save(path)


def png_bytes(img) -> bytes:
    ba = QByteArray()
    buf = QBuffer(ba)
    buf.open(QBuffer.WriteOnly)
    img.save(buf, "PNG")
    return bytes(ba)


class ReportWindow(QMainWindow):
    """Pick the format, the sections and the results' columns, see the report, save it."""

    def __init__(self, main):
        super().__init__(main)
        self.main = main
        self.setWindowTitle("Spikeless — report")
        self.resize(1100, 760)
        side = QWidget()
        side.setFixedWidth(320)
        v = QVBoxLayout(side)
        v.setContentsMargins(0, 0, 0, 0)
        self.c_format = QComboBox()
        self.c_format.addItems([lab for _, lab in FORMATS])
        self.c_format.setToolTip("What Save report writes")
        self.c_format.activated.connect(self._format_set)
        self.c_scope = QComboBox()
        for key, text in SCOPES:
            self.c_scope.addItem(text, key)
        self.c_scope.setToolTip("Several runs: a summary — a line per run, all their results in "
                                "one table — then what was done to each, then their details "
                                "(what they share once, what differs run by run); or each run "
                                "in full in turn. Tick runs in Runs")
        self.c_scope.activated.connect(self._scope_set)
        v.addWidget(QLabel("<b>Runs</b>"))
        v.addWidget(self.c_scope)
        v.addSpacing(8)
        v.addWidget(QLabel("<b>Format</b>"))
        v.addWidget(self.c_format)
        v.addSpacing(8)
        v.addWidget(QLabel("<b>Sections</b>"))
        self.tree = QTreeWidget()
        self.tree.setHeaderHidden(True)
        self.tree.setToolTip("Tick what goes in; the results' columns are under Results (the "
                             "table under the plot has its own: Options › Results)")
        self.tree.itemChanged.connect(self._ticked)
        v.addWidget(self.tree, 1)
        h = QHBoxLayout()
        b = QPushButton("Copy")
        b.setToolTip("The report's text and tables, to paste into a document (the figure has "
                     "its own Copy)")
        b.clicked.connect(self.copy)
        h.addWidget(b)
        b = QPushButton("Save report…")
        b.setToolTip("Ctrl+S")
        b.clicked.connect(self.save)
        h.addWidget(b, 1)
        v.addSpacing(8)
        v.addLayout(h)
        QShortcut(QKeySequence.Save, self, self.save)
        self.view = QTextBrowser()
        root = QWidget()
        hl = QHBoxLayout(root)
        hl.addWidget(side)
        hl.addWidget(self.view, 1)
        self.setCentralWidget(root)
        self.statusBar()
        self._timer = QTimer(self, singleShot=True, interval=250, timeout=self.refresh)
        self.sync()

    def sync(self):
        """Format and ticks from the options (they may have been reset or imported)."""
        exts = [e for e, _ in FORMATS]
        fmt = PREFS["report_format"]
        self.c_format.setCurrentIndex(exts.index(fmt) if fmt in exts else 0)
        self.c_scope.setCurrentIndex(self.c_scope.findData(scope()))
        self.tree.blockSignals(True)
        self.tree.clear()
        for key, label, tip in SECTIONS:
            it = self._tick(self.tree, label, "report_sections", key)
            it.setToolTip(0, tip)
            if key == "results":
                for ck, head, _ in COLUMNS[1:]:          # Region: always there
                    c = self._tick(it, head, "report_cols", ck)
                    c.setToolTip(0, HELP.get(ck, ""))
                it.setExpanded(True)
        self.tree.blockSignals(False)
        self.refresh_later()

    @staticmethod
    def _tick(parent, label, key, sub):
        it = QTreeWidgetItem(parent, [label])
        it.setFlags(Qt.ItemIsEnabled | Qt.ItemIsUserCheckable)
        it.setCheckState(0, Qt.Checked if PREFS[key][sub] else Qt.Unchecked)
        it.setData(0, Qt.UserRole, (key, sub))
        return it

    def _ticked(self, it, _col=0):
        key, sub = it.data(0, Qt.UserRole)
        PREFS[key][sub] = it.checkState(0) == Qt.Checked
        save_prefs()
        self.refresh()

    def _format_set(self, i):
        PREFS["report_format"] = FORMATS[i][0]
        save_prefs()

    def _scope_set(self, i):
        PREFS["report_scope"] = SCOPES[i][0]
        save_prefs()
        self.refresh()

    def open(self):
        self.show()
        self.raise_()
        self.activateWindow()
        self.refresh()

    def refresh_later(self):
        """After a change to the data or the figure: redrawn once things settle, if shown."""
        if self.isVisible():
            self._timer.start()

    def _plot(self):
        """(QImage, width in px at 96 dpi) of the figure as the report shows it, or (None, None)."""
        if not (PREFS["report_sections"]["plot"] and self.main.runs):
            return None, None
        from .figure import figure, render
        fig = figure(self.main)
        return render(fig, 300), round(fig.get_figwidth() * 96)

    def refresh(self):
        img, w = self._plot()
        if img is not None:
            self.view.document().addResource(QTextDocument.ImageResource, QUrl(PLOT), img)
        sb = self.view.verticalScrollBar()
        pos = sb.value()
        self.view.setHtml(html(self.main, w))
        sb.setValue(min(pos, sb.maximum()))

    def copy(self):
        text = html(self.main)
        doc = QTextDocument()
        doc.setHtml(text)
        md = QMimeData()
        md.setHtml(text)
        md.setText(doc.toMarkdown())
        QApplication.clipboard().setMimeData(md)
        self.statusBar().showMessage("Copied (the figure has its own Copy)", 5000)

    def write(self, path: Path) -> None:
        ext = path.suffix.lower()
        img, w = self._plot()
        if ext == ".xlsx":
            return to_xlsx(self.main, path, png_bytes(img) if img else None, w)
        text = html(self.main, w)
        if ext == ".html":
            if img is not None:
                b64 = base64.b64encode(png_bytes(img)).decode("ascii")
                text = text.replace(f"src='{PLOT}'", f"src='data:image/png;base64,{b64}'")
            path.write_text(f"<!doctype html><html><head><meta charset='utf-8'><title>"
                            f"Radio-HPLC report</title></head><body>{text}</body></html>",
                            encoding="utf-8")
            return
        doc = QTextDocument()
        if img is not None:
            doc.addResource(QTextDocument.ImageResource, QUrl(PLOT), img)
        doc.setHtml(text)
        if ext == ".pdf":
            pdf = QPdfWriter(str(path))
            pdf.setPageSize(QPageSize(QPageSize.A4))
            doc.print_(pdf)
        elif ext == ".md":
            md = doc.toMarkdown()
            if img is not None:                       # the figure beside it, as <name>_plot.png
                beside = path.with_name(f"{path.stem}_plot.png")
                beside.write_bytes(png_bytes(img))
                md = md.replace(f"]({PLOT})", f"]({beside.name})")
            path.write_text(md, encoding="utf-8")
        elif not QTextDocumentWriter(str(path), b"ODF").write(doc):
            raise OSError(f"could not write {path}")

    def save(self):
        ext, lab = FORMATS[self.c_format.currentIndex()]
        m = self.main
        stem = m._session_path.stem if m._session_path else \
            (m.runs[0].name if m.runs else "spikeless")
        stem = "".join("_" if c in '\\/:*?"<>|' else c for c in stem)
        path, _ = QFileDialog.getSaveFileName(self, "Save report",
                                              str(m._dir / f"{stem}_report{ext}"),
                                              f"{lab} (*{ext})")
        if not path:
            return
        path = Path(path)
        if not path.name.lower().endswith(ext):       # not with_suffix: names carry dots
            path = path.with_name(path.name + ext)
        try:
            self.write(path)
        except OSError as e:
            QMessageBox.warning(self, "Spikeless", f"Could not write the report:\n{e}")
            return
        self.statusBar().showMessage(f"Written to {path}", 6000)
        m.log(f"Report saved to {path}")
